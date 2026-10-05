#!/usr/bin/env python3
"""PDF2Excel: local PDF/scan to editable Excel converter (GUI + CLI)."""
from __future__ import annotations
import argparse, dataclasses, datetime as dt, gc, logging, math, os, re, shutil, sys, threading, time, traceback
from pathlib import Path
from typing import Any, Callable, Iterable, Optional

APP_NAME = "PDF2Excel"
BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR, OUTPUT_DIR, MODEL_DIR, TEMP_DIR = [BASE_DIR / n for n in ("input", "output", "models", "temp")]
LOG_FILE = BASE_DIR / "pdf2excel.log"
for folder in (INPUT_DIR, OUTPUT_DIR, MODEL_DIR, TEMP_DIR): folder.mkdir(parents=True, exist_ok=True)
logging.basicConfig(filename=LOG_FILE, level=logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s", encoding="utf-8")

class UserError(RuntimeError): pass
class Cancelled(UserError): pass

@dataclasses.dataclass
class Config:
    mode: str = "auto"
    dpi: int = 250
    language: str = "id,en"
    deskew: bool = True
    preprocess: bool = True
    confidence: float = 0.55
    auto_width: bool = True
    merge_cells: bool = True
    preserve_layout: bool = False

@dataclasses.dataclass
class TextBox:
    text: str; x: float; y: float; w: float; h: float; confidence: float = 1.0
    @property
    def cx(self): return self.x + self.w / 2
    @property
    def cy(self): return self.y + self.h / 2

@dataclasses.dataclass
class PageResult:
    page: int; width: int; height: int; boxes: list[TextBox]; rows: list[list[TextBox]]
    table_score: float = 0.0; source: str = "ocr"

@dataclasses.dataclass
class Summary:
    output: Path; pages: int; tables: int; elements: int; warnings: int; seconds: float; quality: int

class Dependency:
    @staticmethod
    def check(require_ocr: bool = False) -> dict[str, str]:
        modules = {"fitz":"PyMuPDF", "cv2":"opencv-python", "numpy":"numpy", "openpyxl":"openpyxl", "PIL":"Pillow"}
        if require_ocr: modules["paddleocr"] = "paddleocr"
        missing = {}
        for mod, package in modules.items():
            try: __import__(mod)
            except Exception as exc: missing[package] = str(exc)
        return missing

class ImageProcessor:
    def __init__(self, config: Config): self.config = config
    def preprocess(self, image):
        import cv2, numpy as np
        gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY) if image.ndim == 3 else image.copy()
        if not self.config.preprocess: return gray
        # CLAHE is robust against the uneven illumination common in office scanners.
        gray = cv2.createCLAHE(2.0, (8, 8)).apply(gray)
        gray = cv2.fastNlMeansDenoising(gray, None, 8, 7, 21)
        if self.config.deskew:
            probe = cv2.adaptiveThreshold(gray,255,cv2.ADAPTIVE_THRESH_GAUSSIAN_C,cv2.THRESH_BINARY_INV,31,15)
            coords = np.column_stack(np.where(probe > 0))
            if len(coords) > 100:
                angle = cv2.minAreaRect(coords[:, ::-1].astype("float32"))[-1]
                angle = -(90 + angle) if angle < -45 else -angle
                if 0.15 < abs(angle) < 15:
                    h,w = gray.shape; matrix = cv2.getRotationMatrix2D((w/2,h/2), angle, 1.0)
                    gray = cv2.warpAffine(gray,matrix,(w,h),flags=cv2.INTER_CUBIC,borderMode=cv2.BORDER_REPLICATE)
        return gray

class OCREngine:
    def __init__(self, config: Config): self.config, self._ocr = config, None
    def _load(self):
        if self._ocr is not None: return
        try:
            from paddleocr import PaddleOCR
            # PaddleOCR 3.x downloads models once to the user's cache; later runs are local.
            try: self._ocr = PaddleOCR(lang="en", use_doc_orientation_classify=True, use_doc_unwarping=False, use_textline_orientation=True)
            except TypeError: self._ocr = PaddleOCR(lang="en", use_angle_cls=True, show_log=False)
        except Exception as exc:
            raise UserError("OCR engine tidak tersedia. Jalankan: pip install -r requirements.txt\nDetail: " + str(exc))
    def recognize(self, image) -> list[TextBox]:
        self._load(); h,w = image.shape[:2]; boxes=[]
        try:
            if hasattr(self._ocr, "predict"):
                results = list(self._ocr.predict(image))
                for result in results:
                    data = getattr(result, "json", result)
                    if callable(data): data=data()
                    if isinstance(data, str):
                        import json; data=json.loads(data)
                    data = data.get("res", data) if isinstance(data, dict) else {}
                    texts=data.get("rec_texts",[]); scores=data.get("rec_scores",[]); polys=data.get("dt_polys",data.get("rec_polys",[]))
                    for text,score,poly in zip(texts,scores,polys):
                        xs=[p[0] for p in poly]; ys=[p[1] for p in poly]
                        boxes.append(TextBox(str(text).strip(),min(xs)/w,min(ys)/h,(max(xs)-min(xs))/w,(max(ys)-min(ys))/h,float(score)))
            else:
                results=self._ocr.ocr(image, cls=True) or []
                for line in results:
                    for item in (line or []):
                        poly,(text,score)=item; xs=[p[0] for p in poly]; ys=[p[1] for p in poly]
                        boxes.append(TextBox(str(text).strip(),min(xs)/w,min(ys)/h,(max(xs)-min(xs))/w,(max(ys)-min(ys))/h,float(score)))
        except Exception as exc: raise UserError("OCR gagal: " + str(exc))
        return [b for b in boxes if b.text]

class TableDetector:
    @staticmethod
    def cluster_rows(boxes: list[TextBox]) -> list[list[TextBox]]:
        if not boxes: return []
        median_h=sorted(b.h for b in boxes)[len(boxes)//2]; tolerance=max(0.004, median_h*0.65)
        rows=[]
        for box in sorted(boxes,key=lambda b:(b.cy,b.x)):
            target=min(rows,key=lambda r:abs(sum(x.cy for x in r)/len(r)-box.cy),default=None)
            if target and abs(sum(x.cy for x in target)/len(target)-box.cy)<=tolerance: target.append(box)
            else: rows.append([box])
        for row in rows: row.sort(key=lambda b:b.x)
        return sorted(rows,key=lambda r:sum(x.cy for x in r)/len(r))
    @staticmethod
    def score(rows):
        useful=[r for r in rows if len(r)>=2]
        return min(1.0,(len(useful)/max(1,len(rows)))*0.65 + (sum(len(r) for r in useful)/max(1,len(useful))/6)*0.35)
    @staticmethod
    def columns(rows):
        xs=sorted(b.x for r in rows for b in r)
        if not xs:return []
        anchors=[]
        for x in xs:
            nearest=min(range(len(anchors)),key=lambda i:abs(anchors[i]-x),default=None)
            if nearest is not None and abs(anchors[nearest]-x)<0.035: anchors[nearest]=(anchors[nearest]+x)/2
            else: anchors.append(x)
        return sorted(anchors)
    def map_rows(self, rows):
        anchors=self.columns(rows); output=[]
        for row in rows:
            cells=[""]*len(anchors)
            for box in row:
                idx=min(range(len(anchors)),key=lambda i:abs(anchors[i]-box.x)); cells[idx]=(cells[idx]+" "+box.text).strip()
            output.append(cells)
        return output

class PDFProcessor:
    def __init__(self, config: Config):
        self.config=config; self.images=ImageProcessor(config); self.ocr=OCREngine(config); self.tables=TableDetector()
    def process(self, pdf_path: Path, progress: Callable[[int,int,str],None], cancel: threading.Event) -> list[PageResult]:
        import fitz, cv2, numpy as np
        try: doc=fitz.open(pdf_path)
        except Exception as exc: raise UserError("PDF tidak dapat dibuka. File mungkin rusak atau terenkripsi.\n"+str(exc))
        if doc.needs_pass: doc.close(); raise UserError("PDF dilindungi password. Buka proteksi lalu coba lagi.")
        if doc.page_count==0: doc.close(); raise UserError("PDF tidak memiliki halaman.")
        results=[]
        try:
            for index in range(doc.page_count):
                if cancel.is_set(): raise Cancelled("Processing cancelled.")
                page=doc[index]; progress(index,doc.page_count,"Menganalisis halaman...")
                digital=[]
                for x0,y0,x1,y1,text,*_ in page.get_text("words"):
                    digital.append(TextBox(text,x0/page.rect.width,y0/page.rect.height,(x1-x0)/page.rect.width,(y1-y0)/page.rect.height,1.0))
                use_ocr=self.config.mode in ("ocr","layout") or len(digital)<8
                if use_ocr:
                    progress(index,doc.page_count,"Rendering dan preprocessing...")
                    pix=page.get_pixmap(matrix=fitz.Matrix(self.config.dpi/72,self.config.dpi/72),alpha=False)
                    image=np.frombuffer(pix.samples,np.uint8).reshape(pix.height,pix.width,pix.n)
                    image=cv2.cvtColor(image,cv2.COLOR_RGB2BGR); processed=self.images.preprocess(image)
                    progress(index,doc.page_count,"Menjalankan OCR..."); boxes=self.ocr.recognize(processed); source="ocr"
                    del image,processed,pix; gc.collect()
                else: boxes=digital; source="digital"
                progress(index,doc.page_count,"Mendeteksi struktur tabel...")
                rows=self.tables.cluster_rows(boxes); score=self.tables.score(rows)
                results.append(PageResult(index+1,int(page.rect.width),int(page.rect.height),boxes,rows,score,source))
                progress(index+1,doc.page_count,"Halaman selesai")
        finally: doc.close()
        return results

class ExcelExporter:
    def __init__(self, config: Config): self.config=config; self.detector=TableDetector()
    @staticmethod
    def value(text: str):
        s=text.strip()
        if re.fullmatch(r"0\d+",s): return s
        if re.fullmatch(r"-?\d+",s):
            try:return int(s)
            except:return s
        if re.fullmatch(r"-?\d+[.,]\d+",s):
            try:return float(s.replace(".","").replace(",","."))
            except:return s
        if re.fullmatch(r"\d+(?:[.,]\d+)?%",s):
            try:return float(s[:-1].replace(",","."))/100
            except:return s
        for fmt in ("%d/%m/%Y","%Y-%m-%d"):
            try:return dt.datetime.strptime(s,fmt).date()
            except ValueError: pass
        return s
    def export(self, pages: list[PageResult], output: Path) -> tuple[int,int]:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.comments import Comment
        from openpyxl.utils import get_column_letter
        wb=Workbook(); wb.remove(wb.active); warnings=0; tables=0
        similar=all(abs(p.table_score-pages[0].table_score)<0.25 for p in pages) if pages else False
        groups=[pages] if similar and sum(p.table_score for p in pages)/len(pages)>0.35 else [[p] for p in pages]
        for gi,group in enumerate(groups,1):
            ws=wb.create_sheet("Data" if len(groups)==1 else f"Page_{group[0].page}")
            row_cursor=1
            for p in group:
                if self.config.mode=="layout" or self.config.preserve_layout:
                    for box in p.boxes:
                        r=max(1,round(box.y*60)+1); c=max(1,round(box.x*20)+1); cell=ws.cell(r,c,self.value(box.text))
                        cell.alignment=Alignment(vertical="top",wrap_text=True)
                        if box.confidence<self.config.confidence: cell.fill=PatternFill("solid",fgColor="FFF2CC"); cell.comment=Comment(f"Low OCR confidence: {box.confidence:.1%}",APP_NAME); warnings+=1
                    row_cursor=max(row_cursor,62)
                else:
                    matrix=self.detector.map_rows(p.rows)
                    if p.table_score>0.35: tables+=1
                    if len(group)>1 and row_cursor>1 and matrix: matrix=matrix[1:] # avoid repeated scan headers
                    for values in matrix:
                        for col,text in enumerate(values,1):
                            cell=ws.cell(row_cursor,col,self.value(text)); cell.alignment=Alignment(vertical="top",wrap_text=True)
                            candidates=[b for b in p.boxes if b.text and b.text in text]
                            conf=min((b.confidence for b in candidates),default=1)
                            if conf<self.config.confidence: cell.fill=PatternFill("solid",fgColor="FFF2CC"); cell.comment=Comment(f"Low OCR confidence: {conf:.1%}",APP_NAME); warnings+=1
                        row_cursor+=1
            if ws.max_row:
                for cell in ws[1]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="205493")
                thin=Side(style="thin",color="B7C9E2")
                for row in ws.iter_rows():
                    for cell in row: cell.border=Border(left=thin,right=thin,top=thin,bottom=thin)
                ws.freeze_panes="A2"
                if ws.max_column>1 and ws.max_row>1: ws.auto_filter.ref=ws.dimensions
                if self.config.auto_width:
                    for col in range(1,ws.max_column+1):
                        width=min(60,max(10,max((len(str(ws.cell(r,col).value or "")) for r in range(1,ws.max_row+1)),default=10)+2))
                        ws.column_dimensions[get_column_letter(col)].width=width
        try: wb.save(output)
        except PermissionError: raise UserError("File Excel sedang terbuka atau folder output tidak dapat ditulis.")
        return tables,warnings

class ConversionEngine:
    def __init__(self, config: Config): self.config=config
    @staticmethod
    def unique_output(path: Path):
        if not path.exists(): return path
        for i in range(1,10000):
            candidate=path.with_name(f"{path.stem}_{i}{path.suffix}")
            if not candidate.exists(): return candidate
        raise UserError("Tidak dapat menentukan nama output unik.")
    def convert(self,input_pdf: Path,output: Optional[Path]=None,progress=lambda a,b,c:None,cancel=None):
        started=time.perf_counter(); cancel=cancel or threading.Event()
        if not input_pdf.exists() or input_pdf.suffix.lower()!=".pdf": raise UserError("Pilih file PDF yang valid.")
        output=self.unique_output(output or OUTPUT_DIR/(input_pdf.stem+".xlsx")); output.parent.mkdir(parents=True,exist_ok=True)
        logging.info("START file=%s mode=%s",input_pdf,self.config.mode)
        pages=PDFProcessor(self.config).process(input_pdf,progress,cancel)
        progress(len(pages),len(pages),"Membangun Excel..."); tables,warnings=ExcelExporter(self.config).export(pages,output)
        elements=sum(len(p.boxes) for p in pages); conf=[b.confidence for p in pages for b in p.boxes]
        quality=round(100*(sum(conf)/len(conf) if conf else 0)*0.75 + 25*(sum(p.table_score for p in pages)/len(pages) if pages else 0))
        summary=Summary(output,len(pages),tables,elements,warnings,time.perf_counter()-started,max(0,min(100,quality)))
        logging.info("DONE output=%s pages=%s elements=%s warnings=%s seconds=%.2f",output,len(pages),elements,warnings,summary.seconds)
        return summary

class PDF2ExcelApp:
    def __init__(self):
        import tkinter as tk
        from tkinter import ttk
        self.tk,self.ttk=tk,ttk; self.root=tk.Tk(); self.root.title(APP_NAME); self.root.geometry("760x620"); self.root.minsize(680,560)
        self.files=[]; self.cancel_event=threading.Event(); self.mode=tk.StringVar(value="auto"); self.status=tk.StringVar(value="Pilih satu atau beberapa PDF"); self.output=tk.StringVar()
        style=ttk.Style(); style.theme_use("vista" if "vista" in style.theme_names() else "clam")
        style.configure("Hero.TLabel",font=("Segoe UI",24,"bold"),foreground="#143A5A"); style.configure("Sub.TLabel",font=("Segoe UI",11),foreground="#52697A")
        style.configure("Primary.TButton",font=("Segoe UI",11,"bold"),padding=12)
        frame=ttk.Frame(self.root,padding=26); frame.pack(fill="both",expand=True)
        ttk.Label(frame,text="PDF SCAN → EXCEL",style="Hero.TLabel").pack(); ttk.Label(frame,text="Convert scanned documents into editable Excel",style="Sub.TLabel").pack(pady=(2,20))
        drop=ttk.LabelFrame(frame,text=" Input PDF ",padding=22); drop.pack(fill="x")
        ttk.Label(drop,text="Pilih PDF (mendukung batch)",font=("Segoe UI",12,"bold")).pack(pady=6)
        ttk.Button(drop,text="SELECT PDF",command=self.select_files,style="Primary.TButton").pack(pady=8)
        self.file_label=ttk.Label(drop,text="Belum ada file",foreground="#52697A"); self.file_label.pack()
        options=ttk.LabelFrame(frame,text=" Conversion Mode ",padding=12); options.pack(fill="x",pady=14)
        for text,value in (("Auto","auto"),("Table","table"),("OCR","ocr"),("Layout Preserve","layout")): ttk.Radiobutton(options,text=text,value=value,variable=self.mode).pack(side="left",padx=16)
        advanced=ttk.LabelFrame(frame,text=" Advanced ",padding=10); advanced.pack(fill="x")
        self.dpi=tk.IntVar(value=250); self.conf=tk.DoubleVar(value=.55); self.deskew=tk.BooleanVar(value=True)
        ttk.Label(advanced,text="DPI").grid(row=0,column=0,padx=7); ttk.Combobox(advanced,textvariable=self.dpi,values=(150,200,250,300,350),width=8,state="readonly").grid(row=0,column=1)
        ttk.Label(advanced,text="Low confidence").grid(row=0,column=2,padx=7); ttk.Spinbox(advanced,from_=0.3,to=.9,increment=.05,textvariable=self.conf,width=8).grid(row=0,column=3)
        ttk.Checkbutton(advanced,text="Deskew",variable=self.deskew).grid(row=0,column=4,padx=14)
        actions=ttk.Frame(frame); actions.pack(pady=18)
        self.convert_btn=ttk.Button(actions,text="CONVERT TO EXCEL",command=self.start,style="Primary.TButton"); self.convert_btn.pack(side="left",padx=6)
        self.cancel_btn=ttk.Button(actions,text="CANCEL",command=self.cancel,state="disabled"); self.cancel_btn.pack(side="left",padx=6)
        self.bar=ttk.Progressbar(frame,mode="determinate"); self.bar.pack(fill="x"); ttk.Label(frame,textvariable=self.status).pack(pady=8)
        self.open_btn=ttk.Button(frame,text="OPEN OUTPUT",command=self.open_output,state="disabled"); self.open_btn.pack()
    def select_files(self):
        from tkinter import filedialog
        names=filedialog.askopenfilenames(title="Select PDF",filetypes=[("PDF files","*.pdf")])
        if names: self.files=[Path(n) for n in names]; self.file_label.config(text=f"{len(self.files)} file: "+", ".join(p.name for p in self.files[:3]))
    def start(self):
        if not self.files:
            from tkinter import messagebox; messagebox.showwarning(APP_NAME,"Pilih PDF terlebih dahulu."); return
        self.cancel_event.clear(); self.convert_btn.config(state="disabled"); self.cancel_btn.config(state="normal"); self.open_btn.config(state="disabled")
        threading.Thread(target=self.worker,daemon=True).start()
    def worker(self):
        try:
            outputs=[]
            for fi,path in enumerate(self.files):
                cfg=Config(mode=self.mode.get(),dpi=self.dpi.get(),confidence=self.conf.get(),deskew=self.deskew.get(),preserve_layout=self.mode.get()=="layout")
                def progress(done,total,msg): self.root.after(0,self.update_progress,fi,done,total,msg)
                summary=ConversionEngine(cfg).convert(path,progress=progress,cancel=self.cancel_event); outputs.append(summary)
            self.output.set(str(outputs[-1].output)); total_pages=sum(s.pages for s in outputs); warnings=sum(s.warnings for s in outputs)
            self.root.after(0,self.finish,f"Selesai: {len(outputs)} file, {total_pages} halaman, {warnings} warning. Quality indicator: {outputs[-1].quality}%")
        except Exception as exc:
            logging.error("%s\n%s",exc,traceback.format_exc()); self.root.after(0,self.fail,str(exc))
    def update_progress(self,fi,done,total,msg):
        pct=((fi+done/max(1,total))/len(self.files))*100; self.bar["value"]=pct; self.status.set(f"File {fi+1}/{len(self.files)} · Page {done}/{total} · {msg} · {pct:.0f}%")
    def finish(self,msg): self.bar["value"]=100; self.status.set(msg); self.convert_btn.config(state="normal"); self.cancel_btn.config(state="disabled"); self.open_btn.config(state="normal")
    def fail(self,msg):
        from tkinter import messagebox
        self.convert_btn.config(state="normal"); self.cancel_btn.config(state="disabled"); self.status.set("Konversi gagal"); messagebox.showerror("Cannot convert PDF",msg)
    def cancel(self): self.cancel_event.set(); self.status.set("Membatalkan...")
    def open_output(self):
        path=Path(self.output.get())
        if path.exists(): os.startfile(path) if os.name=="nt" else os.system(f'xdg-open "{path}"')
    def run(self): self.root.mainloop()

def cli_progress(done,total,msg):
    pct=round(done/max(1,total)*100); width=24; bar="█"*round(width*pct/100)+"░"*(width-round(width*pct/100))
    print(f"\r[{bar}] {pct:3d}%  Page {done}/{total}  {msg:<28}",end="",flush=True)

def main(argv=None):
    parser=argparse.ArgumentParser(prog="pdf_to_excel.py",description="Convert PDF/scan into editable Excel.")
    parser.add_argument("input",nargs="?",type=Path); parser.add_argument("-o","--output",type=Path); parser.add_argument("--mode",choices=("auto","table","ocr","layout"),default="auto")
    parser.add_argument("--dpi",type=int,choices=(150,200,250,300,350),default=250); parser.add_argument("--confidence",type=float,default=.55); parser.add_argument("--diagnose",action="store_true")
    args=parser.parse_args(argv)
    missing=Dependency.check(require_ocr=bool(args.input and args.mode in ("ocr","layout")))
    if args.diagnose:
        print("PDF2Excel diagnostic"); print("Python:",sys.version.split()[0]); print("Writable output:",os.access(OUTPUT_DIR,os.W_OK)); print("Missing:",", ".join(missing) or "none"); return 1 if missing else 0
    if not args.input:
        basic=Dependency.check(False)
        if basic: print("Missing dependency: "+", ".join(basic)+"\nRun: pip install -r requirements.txt"); return 2
        try: PDF2ExcelApp().run(); return 0
        except Exception as exc: print("GUI tidak dapat dibuka:",exc); logging.error(traceback.format_exc()); return 2
    try:
        print("PDF2Excel\n"+"─"*30+f"\nInput : {args.input}\nMode  : {args.mode.upper()}\n")
        cfg=Config(mode=args.mode,dpi=args.dpi,confidence=args.confidence,preserve_layout=args.mode=="layout")
        summary=ConversionEngine(cfg).convert(args.input,args.output,cli_progress); print(f"\n\nConversion completed.\nOutput: {summary.output}\nPages: {summary.pages} | Tables: {summary.tables} | OCR elements: {summary.elements} | Warnings: {summary.warnings}\nQuality indicator: {summary.quality}% (review recommended) | Time: {summary.seconds:.1f}s")
        return 0
    except Exception as exc: logging.error("%s\n%s",exc,traceback.format_exc()); print(f"\nCannot convert PDF\nReason: {exc}\nDetail teknis: {LOG_FILE}"); return 1

if __name__=="__main__": raise SystemExit(main())
